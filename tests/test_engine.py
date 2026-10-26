"""Run with:  python -m pytest -q"""
import datetime as dt
import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from engine import inspect_json, process
from engine.demo_data import generate
from engine.loader import lead_bucket, person, phone, pipe_stage

END = dt.date(2026, 10, 25)


@pytest.fixture(scope="module")
def demo():
    return generate(END)


@pytest.fixture(scope="module")
def result(demo):
    return process(demo, "2026-10-26")


def test_detection(demo):
    kinds = {json.loads(inspect_json(n, b))["kind"] for n, b in demo}
    assert kinds == {"meta", "indiamart", "justdial", "pipeline"}


def test_status_mapping():
    assert lead_bucket("QUALIFIED") == "Qualified"
    assert lead_bucket("NOT_QUALIFIED") == "Not Qualified"
    assert lead_bucket("CALL_NOT_PICK") == "Call Not Pick"
    assert lead_bucket("CREATED") == "Pending (Not Worked)"
    assert lead_bucket("Quotation Sent") == "Qualified"
    assert pipe_stage("Qualified") == "No Action"
    assert pipe_stage("Won") == "Won" and pipe_stage("Lost") == "Lost"
    assert pipe_stage("quotation sent") == "Quotation Sent"


def test_cleaning_helpers():
    assert person(" UJALA  RAJPUT ") == "Ujala Rajput"
    assert person(None) == "Unassigned" and person(float("nan")) == "Unassigned"
    assert phone(9829032400.0) == "9829032400"
    assert phone("+91 98290 32400") == "9829032400"


def test_totals_reconcile(result):
    s, _ = result
    k = s["kpi"]
    assert k["leads"] == 150 + 40 + 12
    assert k["qualified"] + k["not_qualified"] + k["call_not_pick"] + k["pending"] + k["other"] == k["leads"]
    assert sum(r["leads"] for r in s["scorecard"]["rows"]) == k["leads"]
    assert sum(r["leads"] for r in s["sources"]["rows"]) == k["leads"]
    assert k["pipe_total"] == 60
    assert k["no_action"] + k["actioned"] == k["pipe_total"]
    assert sum(sum(stg.values()) for stg in [s["pipeline"]["team"]["stages"]]) == k["pipe_total"]
    assert len(s["followup"]) == k["no_action"]
    assert len(s["unassigned"]) == k["unassigned"]


def test_summary_is_json_safe(result):
    s, _ = result
    json.dumps(s, allow_nan=False)


def test_workbook(result):
    _, xlsx = result
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Lead Report", "Data_Leads", "Data_Pipeline"]
    assert wb["Data_Leads"].max_row == 150 + 40 + 12 + 1
    assert str(wb["Lead Report"]["A7"].value).startswith("=")


def test_partial_upload(demo):
    s, xlsx = process([demo[0]], "2026-10-26")
    assert s["kpi"]["leads"] == 150 and s["kpi"]["pipe_total"] == 0 and xlsx
    s, xlsx = process([demo[3]], "2026-10-26")
    assert s["kpi"]["leads"] == 0 and s["kpi"]["pipe_total"] == 60 and xlsx


def test_unrecognised_file():
    buf = io.BytesIO(); pd.DataFrame({"foo": [1]}).to_excel(buf, index=False)
    s, xlsx = process([("random.xlsx", buf.getvalue())], "2026-10-26")
    assert xlsx is None
    assert any("not recognised" in q["text"] for q in s["quality"])
